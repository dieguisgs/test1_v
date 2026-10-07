"""Exercise per-product settings through actual CSV/XLSX production commands."""

import json
from dataclasses import replace
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from vwaps import cli
from vwaps.config import load_config
from vwaps.enrich import enrich_input
from vwaps.mapping import COLUMNS, ProductMap, read_mapping_table
from vwaps.product_config import MODEL_PARAMETER_FIELDS
from vwaps.publication import CsvBatch


DAY = date(2026, 9, 29)


def project(root, extension="csv"):
    root.mkdir(parents=True, exist_ok=True)
    config = root / "config.toml"
    config.write_text(
        '[paths]\nvwap_input="own.csv"\nmapping="mapping.' + extension + '"\n'
        'eex_curves_dir="."\noutput_dir="output"\n'
        '[run]\nconfiguration_mode="individual"\n'
        '[targets]\ntenors=["M+1", "M+2"]\n'
        '[layers]\nhist="off"\n[method]\nbasis_mode="auto"\n'
        '[vwap_columns]\ntenor="tenor"\nvolume="volume"\n', encoding="utf-8")
    table = pd.DataFrame([
        dict(product="POWER", region=region, unit="EUR/MWh", use="fill", area="DE",
             profile="Base", eex_file="eex.csv", hours="Base", timezone="Europe/Berlin",
             comment="literal = note", basis_mode=mode, tau_log=tau,
             tenors=json.dumps(targets))
        for region, mode, tau, targets in (
            ("001", "additive", "0.1", ["M+1", "M+2"]),
            ("002", "ratio", "2.0", ["M+1", "M+2", "M+3"]))
    ], columns=[*COLUMNS, "basis_mode", "tau_log", "tenors"])
    path = root / f"mapping.{extension}"
    if extension == "xlsx":
        table.to_excel(path, index=False)
    else:
        table.to_csv(path, index=False)
    pd.DataFrame([
        dict(reference_date=str(DAY), product="POWER", region=region,
             unit="EUR/MWh", tenor="M+1", vwap=110., volume=100.)
        for region in ("001", "002")
    ]).to_csv(root / "own.csv", index=False)
    pd.DataFrame([
        dict(tradeDate=str(DAY), maturityType="Month", deliveryStart=delivery, settlPx=price)
        for delivery, price in (("2026-10-01", 100.), ("2026-11-01", 200.), ("2026-12-01", 300.))
    ]).to_csv(root / "eex.csv", index=False)
    return config


def filled(config):
    return pd.read_csv(config.parent / "output" / "filled_history.csv", dtype=str, keep_default_na=False)


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_daily_refill_catchup_use_same_individual_configuration(tmp_path, extension):
    frames = {}
    for verb in ("daily", "refill", "catchup"):
        config = project(tmp_path / verb, extension)
        args = [verb, "--date", str(DAY)] if verb == "daily" else [verb, "--from", str(DAY), "--to", str(DAY)]
        assert cli.main(args, config) == 0
        frames[verb] = filled(config)
    pd.testing.assert_frame_equal(frames["daily"], frames["refill"])
    pd.testing.assert_frame_equal(frames["daily"], frames["catchup"])
    frame = frames["daily"]
    assert frame.groupby("region").size().to_dict() == {"001": 2, "002": 3}
    assert set(frame.configuration_mode) == {"individual"}
    assert frame.configuration_id.nunique() == 2
    parameters = {row.region: json.loads(row.configuration_parameters) for row in frame.itertuples()}
    assert parameters["001"]["tau_log"] == .1
    assert parameters["002"]["basis_mode"] == "ratio"


def test_global_cli_ignores_invalid_individual_cells(tmp_path):
    config = project(tmp_path)
    table = read_mapping_table(tmp_path / "mapping.csv")
    table["tau_log"] = "invalid and deliberately ignored"
    table.to_csv(tmp_path / "mapping.csv", index=False)
    assert cli.main(["daily", "--date", str(DAY), "--configuration-mode", "global"], config) == 0
    frame = filled(config)
    assert set(frame.configuration_mode) == {"global"}
    assert len(frame) == 4
    assert frame.configuration_id.nunique() == 1
    assert cli.main(["daily", "--date", str(DAY)], config) == 1


def test_explicit_cli_flags_override_individual_cells(tmp_path):
    config = project(tmp_path)
    table = read_mapping_table(tmp_path / "mapping.csv")
    table["shape_mode"] = "adjust"
    table["shape_adjust_originals"] = "true"
    table.to_csv(tmp_path / "mapping.csv", index=False)
    assert cli.main(["daily", "--date", str(DAY), "--shape-mode", "off",
                     "--shape-adjust-originals", "off"], config) == 0
    for payload in filled(config).configuration_parameters:
        parameters = json.loads(payload)
        assert parameters["shape_mode"] == "off"
        assert parameters["shape_adjust_originals"] is False


def test_catchup_rejects_changed_settings_before_output_mutation(tmp_path):
    config = project(tmp_path)
    args = ["catchup", "--from", str(DAY), "--to", str(DAY)]
    assert cli.main(args, config) == 0
    table = read_mapping_table(tmp_path / "mapping.csv")
    table.loc[table.region.eq("001"), "tau_log"] = "0.8"
    table.to_csv(tmp_path / "mapping.csv", index=False)
    snapshots = {path: path.read_bytes() for path in (tmp_path / "output").rglob("*.csv")}
    assert cli.main(args, config) == 1
    assert {path: path.read_bytes() for path in snapshots} == snapshots
    assert cli.main(["refill", "--from", str(DAY), "--to", str(DAY)], config) == 0
    assert cli.main(args, config) == 0


def test_individual_target_expansion_recalculates_only_affected_curve(tmp_path):
    config = project(tmp_path)
    args = ["catchup", "--from", str(DAY), "--to", str(DAY)]
    assert cli.main(args, config) == 0
    before = filled(config).query("region == '002'").reset_index(drop=True)
    table = read_mapping_table(tmp_path / "mapping.csv")
    table.loc[table.region.eq("001"), "tenors"] = '["M+1", "M+2", "M+3"]'
    table.to_csv(tmp_path / "mapping.csv", index=False)
    assert cli.main(args, config) == 0
    after = filled(config)
    assert len(after.query("region == '001'")) == 3
    pd.testing.assert_frame_equal(after.query("region == '002'").reset_index(drop=True), before)


def test_mapping_adds_columns_and_exports_excel_without_losing_values(tmp_path):
    config = project(tmp_path)
    export = tmp_path / "editable.xlsx"
    assert cli.main(["mapping", "--with-parameters", "--export", str(export)], config) == 0
    table = read_mapping_table(export)
    assert all(name in table for name in MODEL_PARAMETER_FIELDS)
    assert list(table.region) == ["001", "002"]
    assert list(table.tau_log) == ["0.1", "2.0"]
    snapshot = export.read_bytes()
    assert cli.main(["mapping", "--export", str(export)], config) == 1
    assert export.read_bytes() == snapshot


def test_mapping_excel_update_preserves_other_worksheets(tmp_path):
    config = project(tmp_path, "xlsx")
    path = tmp_path / "mapping.xlsx"
    book = load_workbook(path)
    notes = book.create_sheet("Notes")
    notes["A1"] = "Keep this note"
    book.save(path)
    book.close()
    assert cli.main(["mapping", "--with-parameters"], config) == 0
    updated = load_workbook(path)
    assert updated["Notes"]["A1"].value == "Keep this note"
    assert updated.worksheets[0].freeze_panes == "D2"
    updated.close()
    assert "shape_mode" in read_mapping_table(path)


def test_excel_publication_does_not_turn_literals_into_formulas(tmp_path):
    path = tmp_path / "mapping.xlsx"
    with CsvBatch() as batch:
        batch.stage_excel(path, pd.DataFrame({"comment": ["=literal", None]}))
    book = load_workbook(path)
    assert book.active["A2"].data_type == "s"
    assert book.active["A2"].value == "=literal"
    assert book.active["A3"].value is None
    book.close()


def test_enrichment_uses_individual_permission_to_adjust_originals(tmp_path):
    config = project(tmp_path)
    cfg = load_config(config)
    mapping = ProductMap("POWER", "fill", "DE", "Base", "eex.csv", "Base", "Europe/Berlin",
                         region="001", unit="EUR/MWh",
                         parameter_overrides={"shape_mode": "adjust", "shape_adjust_originals": True})
    raw = pd.DataFrame([dict(reference_date=str(DAY), product="POWER", region="001", unit="EUR/MWh",
                             tenor="M+1", vwap=110., volume=100.)])
    result = pd.DataFrame([dict(reference_date=DAY, product="POWER", region="001", unit="EUR/MWh",
                                tenor="M+1", price=105., source="own+shape", price_before_shape=110.,
                                shape_proposed_adjustment=-5., shape_adjustment=-5., shape_original_modified=True)])
    enriched = enrich_input(raw, result, cfg, [mapping], DAY, DAY)
    assert enriched.iloc[0].vwap == 110.
    assert enriched.iloc[0].curve_price == 105.
    assert enriched.iloc[0].data_origin == "estimated"
    assert enriched.iloc[0].estimation_method == "shape_adjusted_original"


@pytest.mark.parametrize("cross", [False, True])
def test_catchup_checks_helper_configuration_only_when_cross_is_enabled(tmp_path, cross):
    config = project(tmp_path)
    table = read_mapping_table(tmp_path / "mapping.csv")
    table.loc[table.region.eq("002"), "use"] = "helper"
    table["layer_cross"] = "true" if cross else "false"
    table.to_csv(tmp_path / "mapping.csv", index=False)
    args = ["catchup", "--from", str(DAY), "--to", str(DAY)]
    assert cli.main(args, config) == 0
    before = filled(config)
    context_id = before.iloc[0].configuration_context_id
    context_path = tmp_path / "output" / "configurations" / f"{context_id}.json"
    assert context_path.exists()
    context = json.loads(context_path.read_text(encoding="utf-8"))
    assert context["schema_version"] == 1
    table["min_volume"] = ""
    table.loc[table.region.eq("002"), "min_volume"] = "1000"
    table.to_csv(tmp_path / "mapping.csv", index=False)
    assert cli.main(args, config) == (1 if cross else 0)
    if cross:
        assert cli.main(["refill", "--from", str(DAY), "--to", str(DAY)], config) == 0
        after = filled(config)
        assert before.iloc[0].configuration_id == after.iloc[0].configuration_id
        assert before.iloc[0].configuration_context_id != after.iloc[0].configuration_context_id
        assert context_path.exists()  # Older context snapshots remain auditable.


def test_invalid_excel_parameter_cannot_overwrite_existing_output(tmp_path):
    config = project(tmp_path, "xlsx")
    args = ["daily", "--date", str(DAY)]
    assert cli.main(args, config) == 0
    output = tmp_path / "output" / "filled_history.csv"
    before = output.read_bytes()
    table = read_mapping_table(tmp_path / "mapping.xlsx")
    table.loc[0, "tau_log"] = "not a number"
    table.to_excel(tmp_path / "mapping.xlsx", index=False)
    assert cli.main(args, config) == 1
    assert output.read_bytes() == before
